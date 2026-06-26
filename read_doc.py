import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

doc = Document(r'c:\Users\adone\OneDrive\桌面\AI练习\wendang\fangan\【协会反馈】绿色机场平台二期-建设方案V.020260608.docx')

for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f'[{i}] {p.style.name}: {p.text[:200]}')

# Also check tables
print('\n=== TABLES ===')
for ti, table in enumerate(doc.tables):
    print(f'\nTable {ti}:')
    for ri, row in enumerate(table.rows):
        cells = [cell.text.strip()[:50] for cell in row.cells]
        print(f'  Row {ri}: {cells}')
