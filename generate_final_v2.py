import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

# 使用原始文件作为源
source_path = r'c:\Users\adone\OneDrive\桌面\AI练习\wendang\fangan\【协会反馈】绿色机场平台二期-建设方案V.020260608.docx'
doc = Document(source_path)

changes_log = []

def find_para_by_contains(doc, text_fragment):
    for i, p in enumerate(doc.paragraphs):
        if text_fragment in p.text:
            return i
    return -1

def replace_para_text(doc, idx, new_text):
    if idx < 0 or idx >= len(doc.paragraphs):
        return False
    p = doc.paragraphs[idx]
    for run in p.runs:
        run.text = ''
    if p.runs:
        p.runs[0].text = new_text
    else:
        p.add_run(new_text)
    return True

def delete_para(doc, idx):
    if idx < 0 or idx >= len(doc.paragraphs):
        return
    p = doc.paragraphs[idx]
    p._element.getparent().remove(p._element)

# ============================================================
# 修改1: 删除所有"等保三级"相关内容
# ============================================================
# 封面
idx = find_para_by_contains(doc, '阿里云生态部署 · 等保三级合规')
if idx >= 0:
    replace_para_text(doc, idx, '阿里云生态部署')
    changes_log.append('[1] 封面：删除等保三级合规')

# 调整一
idx = find_para_by_contains(doc, '等保三级合规要求')
if idx >= 0:
    text = doc.paragraphs[idx].text
    new_text = text.replace('等保三级合规要求', '相关合规要求')
    replace_para_text(doc, idx, new_text)
    changes_log.append('[2] 删除：等保三级合规要求')

# 落实等保三级合规
idx = find_para_by_contains(doc, '落实等保三级合规')
if idx >= 0:
    text = doc.paragraphs[idx].text
    new_text = text.replace('落实等保三级合规', '落实相关合规')
    replace_para_text(doc, idx, new_text)
    changes_log.append('[3] 删除：落实等保三级合规')

# 系统安全中的等保三级合规
idx = find_para_by_contains(doc, '系统安全：六角色权限体系')
if idx >= 0:
    text = doc.paragraphs[idx].text
    if '等保三级合规' in text:
        new_text = text.replace('+ 等保三级合规', '')
        replace_para_text(doc, idx, new_text)
        changes_log.append('[4] 系统安全：删除等保三级合规')

# 删除7.2 等保三级合规章节
idx = find_para_by_contains(doc, '7.2 等保三级合规')
if idx >= 0:
    delete_para(doc, idx)
    changes_log.append('[5] 删除：7.2 等保三级合规章节')

idx = find_para_by_contains(doc, '阿里云已具备等保三级基础设施层资质')
if idx >= 0:
    delete_para(doc, idx)
    changes_log.append('[6] 删除：阿里云已具备等保三级基础设施层资质')

idx = find_para_by_contains(doc, '物理与环境：依托阿里云数据中心已具备的等保三级资质')
if idx >= 0:
    delete_para(doc, idx)
    changes_log.append('[7] 删除：物理与环境等保三级资质')

# 安全组件描述
idx = find_para_by_contains(doc, '等保三级所需安全组件')
if idx >= 0:
    text = doc.paragraphs[idx].text
    new_text = text.replace('等保三级所需安全组件', '所需安全组件')
    replace_para_text(doc, idx, new_text)
    changes_log.append('[8] 修改：等保三级所需安全组件')

idx = find_para_by_contains(doc, '安全组与等保三级合规组件')
if idx >= 0:
    text = doc.paragraphs[idx].text
    new_text = text.replace('安全组与等保三级合规组件', '安全组与相关合规组件')
    replace_para_text(doc, idx, new_text)
    changes_log.append('[9] 修改：安全组与等保三级合规组件')

# 表格中的等保三级
table_count = 0
for table in doc.tables:
    table_count += 1
    for row in table.rows:
        for cell in row.cells:
            if '等保三级' in cell.text:
                for para in cell.paragraphs:
                    for run in para.runs:
                        if '等保三级' in run.text:
                            run.text = run.text.replace('等保三级', '')
                            changes_log.append(f'[10] 表格{table_count}中删除：等保三级')

# 段落中的剩余等保三级
for i, p in enumerate(doc.paragraphs):
    if '等保三级' in p.text:
        for run in p.runs:
            if '等保三级' in run.text:
                run.text = run.text.replace('等保三级', '')
                changes_log.append(f'[11] 段落{i}删除：等保三级')

# ============================================================
# 修改2: "9类增值服务"改为"7项增值服务"
# ============================================================
# 1.2目标二
idx = find_para_by_contains(doc, '新增 9 类增值服务功能模块')
if idx >= 0:
    text = doc.paragraphs[idx].text
    new_text = text.replace('新增 9 类增值服务功能模块', '新增7项增值服务功能模块')
    replace_para_text(doc, idx, new_text)
    changes_log.append('[12] 1.2目标二：9类→7项')

# 4.1设计动机与定位
idx = find_para_by_contains(doc, '规划 9 项面向会员机场单位的增值服务模块')
if idx >= 0:
    text = doc.paragraphs[idx].text
    new_text = text.replace('规划 9 项面向会员机场单位的增值服务模块', '规划7项面向会员机场单位的增值服务模块')
    replace_para_text(doc, idx, new_text)
    changes_log.append('[13] 4.1：9项→7项')

# 4.11总览中的9项扩展模块
idx = find_para_by_contains(doc, '9 项扩展模块共同实现这一定位转变')
if idx >= 0:
    text = doc.paragraphs[idx].text
    new_text = text.replace('9 项扩展模块共同实现这一定位转变', '7项扩展模块共同实现这一定位转变')
    replace_para_text(doc, idx, new_text)
    changes_log.append('[14] 4.11：9项→7项')

# ============================================================
# 修改3: 修改PDF报告数据抓取方式
# ============================================================
idx = find_para_by_contains(doc, '四星评价强制')
if idx >= 0:
    text = doc.paragraphs[idx].text
    if '图片转文字' not in text and 'OCR' not in text:
        new_text = text + ' 二期新建立数据库，通过图片转文字算法逻辑做抓取，避免AI的直接参与导致数据信息泄露。'
        replace_para_text(doc, idx, new_text)
        changes_log.append('[15] PDF抓取：添加图片转文字算法逻辑')
    elif 'OCR' in text and '图片转文字' not in text:
        new_text = text.replace('OCR（光学字符识别）+ NLP（自然语言处理）做抓取', '图片转文字算法逻辑做抓取')
        replace_para_text(doc, idx, new_text)
        changes_log.append('[15] PDF抓取：OCR+NLP改为图片转文字算法逻辑')

# ============================================================
# 保存最终版本
# ============================================================
output_path = r'c:\Users\adone\OneDrive\桌面\AI练习\wendang\fangan\【协会反馈】绿色机场平台二期-建设方案V.020260608_最终修改版.docx'
doc.save(output_path)

# 输出修改日志
print("=" * 70)
print("文档修改完成！")
print("=" * 70)
for log in changes_log:
    print(log)
print("=" * 70)
print(f"共完成 {len(changes_log)} 项修改")
print(f"\n📄 已保存为：{output_path}")