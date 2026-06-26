import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

# 读取原始文件
source_path = r'c:\Users\adone\OneDrive\桌面\AI练习\wendang\fangan\【协会反馈】绿色机场平台二期-建设方案V.020260608.docx'
source_doc = Document(source_path)

# 读取修改后的文件
modified_path = r'c:\Users\adone\OneDrive\桌面\AI练习\wendang\fangan\【协会反馈】绿色机场平台二期-建设方案V.020260608_最终修改版.docx'
modified_doc = Document(modified_path)

print("=" * 70)
print("原始文件 vs 修改后文件 对比")
print("=" * 70)

# 检查1: 等保内容
print("\n【等保内容检查】")
source_dengbao = [p.text for p in source_doc.paragraphs if '等保' in p.text]
modified_dengbao = [p.text for p in modified_doc.paragraphs if '等保' in p.text]
print(f"原始文件等保段落数: {len(source_dengbao)}")
print(f"修改后文件等保段落数: {len(modified_dengbao)}")
if source_dengbao:
    print("原始文件等保内容:")
    for t in source_dengbao[:3]:
        print(f"  - {t[:60]}")

# 检查2: 增值服务数量
print("\n【增值服务数量检查】")
source_9type = [p.text for p in source_doc.paragraphs if '9 类' in p.text or '9 项' in p.text]
modified_7type = [p.text for p in modified_doc.paragraphs if '7项' in p.text]
modified_9type = [p.text for p in modified_doc.paragraphs if '9 类' in p.text or '9 项' in p.text]
print(f"原始文件'9类/9项'段落: {len(source_9type)}")
print(f"修改后文件'7项'段落: {len(modified_7type)}")
print(f"修改后文件仍有'9类/9项'段落: {len(modified_9type)}")

# 检查3: PDF抓取方式
print("\n【PDF抓取方式检查】")
source_pdf = [p.text for p in source_doc.paragraphs if 'PDF' in p.text or '关键词' in p.text or '抓取' in p.text]
modified_pdf = [p.text for p in modified_doc.paragraphs if '图片转文字' in p.text or '避免AI' in p.text]
print(f"原始文件PDF相关段落: {len(source_pdf)}")
print(f"修改后文件图片转文字相关段落: {len(modified_pdf)}")
for t in modified_pdf:
    print(f"  - {t[:100]}")

# 检查4: 文件大小对比
import os
source_size = os.path.getsize(source_path)
modified_size = os.path.getsize(modified_path)
print(f"\n【文件大小对比】")
print(f"原始文件大小: {source_size} 字节")
print(f"修改后文件大小: {modified_size} 字节")
print(f"差异: {modified_size - source_size} 字节")

# 检查5: 检查修改后的文件是否真的包含修改内容
print("\n【修改后文件关键内容验证】")
all_modified_text = ' '.join([p.text for p in modified_doc.paragraphs])

checks = [
    ('等保', '等保内容', False),
    ('7项增值服务', '7项增值服务', True),
    ('图片转文字', '图片转文字算法', True),
    ('避免AI', '避免AI直接参与', True),
]

for keyword, label, should_exist in checks:
    exists = keyword in all_modified_text
    status = "✅" if exists == should_exist else "❌"
    print(f"{status} {label}: {'存在' if exists else '不存在'}")

print("\n" + "=" * 70)