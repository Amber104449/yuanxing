import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

# 检查原始文档
doc_path = r'c:\Users\adone\OneDrive\桌面\AI练习\wendang\fangan\【协会反馈】绿色机场平台二期-建设方案V.020260608.docx'
doc = Document(doc_path)

# 搜索所有包含"PDF"、"数据库"、"抓取"、"关键词"的段落
for i, p in enumerate(doc.paragraphs):
    if 'PDF' in p.text or '抓取' in p.text or '数据库' in p.text or '关键词' in p.text:
        print(f"段落[{i}]: {p.text}")

# 搜索3.4或方向相关内容
for i, p in enumerate(doc.paragraphs):
    if '方向' in p.text and '八' in p.text:
        print(f"\n段落[{i}]: {p.text}")

# 搜索碳达峰
for i, p in enumerate(doc.paragraphs):
    if '碳达峰' in p.text:
        print(f"\n段落[{i}]: {p.text}")