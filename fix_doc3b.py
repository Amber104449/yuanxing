import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn

doc_path = r'c:\Users\adone\OneDrive\桌面\AI练习\wendang\fangan\【协会反馈】绿色机场平台二期-建设方案V.020260608_修改版.docx'
doc = Document(doc_path)

changes_log = []

def find_para_by_contains(doc, text_fragment):
    for i, p in enumerate(doc.paragraphs):
        if text_fragment in p.text:
            return i
    return -1

def replace_para_text(doc, idx, new_text):
    if idx < 0 or idx >= len(doc.paragraphs):
        return False
    p = doc.paragraphs[idx]
    for run in p.runs:
        run.text = ''
    if p.runs:
        p.runs[0].text = new_text
    else:
        p.add_run(new_text)
    return True

def delete_para(doc, idx):
    if idx < 0 or idx >= len(doc.paragraphs):
        return
    p = doc.paragraphs[idx]
    p._element.getparent().remove(p._element)

def remove_text_in_cell(cell, text_to_remove):
    """从单元格中删除指定文本"""
    for para in cell.paragraphs:
        for run in para.runs:
            if text_to_remove in run.text:
                run.text = run.text.replace(text_to_remove, '')
                return True
    return False

# ============================================================
# 处理表格中的"等保三级"
# ============================================================
table_count = 0
for table in doc.tables:
    table_count += 1
    for row in table.rows:
        for cell in row.cells:
            if '等保三级' in cell.text:
                # 遍历每个段落
                for para in cell.paragraphs:
                    for run in para.runs:
                        if '等保三级' in run.text:
                            run.text = run.text.replace('等保三级', '')
                            changes_log.append(f'[TABLE] 表格{table_count}中删除：等保三级')

# 再次检查段落中是否还有"等保三级"
for i, p in enumerate(doc.paragraphs):
    if '等保三级' in p.text:
        for run in p.runs:
            if '等保三级' in run.text:
                run.text = run.text.replace('等保三级', '')
                changes_log.append(f'[PARA {i}] 删除：等保三级')

# 保存文档
doc.save(doc_path)

# 输出修改日志
print("=" * 50)
print("表格等保三级清理完成！")
print("=" * 50)
for log in changes_log:
    print(log)
print("=" * 50)
print(f"共完成 {len(changes_log)} 项修改")