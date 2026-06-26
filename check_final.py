import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

doc_path = r'c:\Users\adone\OneDrive\桌面\AI练习\wendang\fangan\【协会反馈】绿色机场平台二期-建设方案V.020260608_最终修改版.docx'
doc = Document(doc_path)

# 查找PDF相关内容
for i, p in enumerate(doc.paragraphs):
    if 'PDF' in p.text or '数据库' in p.text or '关键词' in p.text:
        print(f"段落[{i}]: {p.text}")

# 查找绿色机场平台二期相关内容
for i, p in enumerate(doc.paragraphs):
    if '绿色机场平台二期' in p.text or '建设方案' in p.text:
        print(f"段落[{i}]: {p.text}")

# 检查等保
has_dengbao = any('等保' in p.text for p in doc.paragraphs)
print(f"\n等保内容: {'存在' if has_dengbao else '已删除'}")

# 检查增值服务
has_7items = any('7项' in p.text for p in doc.paragraphs)
print(f"7项增值服务: {'存在' if has_7items else '未找到'}")

# 检查图片转文字
has_ocr = any('图片转文字' in p.text for p in doc.paragraphs)
print(f"图片转文字: {'存在' if has_ocr else '未找到'}")