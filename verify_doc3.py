import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

doc_path = r'c:\Users\adone\OneDrive\桌面\AI练习\wendang\fangan\【协会反馈】绿色机场平台二期-建设方案V.020260608_修改版.docx'
doc = Document(doc_path)

all_text = ' '.join([p.text for p in doc.paragraphs])

# 检查等保三级
has_dengbao = '等保三级' in all_text
print(f"文档中是否还有'等保三级': {has_dengbao}")

# 检查7项增值服务
has_7items = '7项增值服务' in all_text
print(f"文档中是否包含'7项增值服务': {has_7items}")

# 检查图片转文字
has_ocr = '图片转文字算法逻辑' in all_text
print(f"文档中是否包含'图片转文字算法逻辑': {has_ocr}")

# 检查避免AI直接参与
has_no_ai = '避免AI的直接参与' in all_text
print(f"文档中是否包含'避免AI的直接参与': {has_no_ai}")

# 搜索增值服务相关内容
for i, p in enumerate(doc.paragraphs):
    if '增值服务' in p.text:
        print(f"[{i}] {p.text[:100]}")