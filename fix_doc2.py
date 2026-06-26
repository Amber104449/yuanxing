import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

doc_path = r'c:\Users\adone\OneDrive\桌面\AI练习\wendang\fangan\【协会反馈】绿色机场平台二期-建设方案V.020260608_修改版.docx'
doc = Document(doc_path)

changes_log = []

def find_para_by_contains(doc, text_fragment):
    for i, p in enumerate(doc.paragraphs):
        if text_fragment in p.text:
            return i
    return -1

def replace_para_text(doc, idx, new_text):
    if idx < 0 or idx >= len(doc.paragraphs):
        return False
    p = doc.paragraphs[idx]
    for run in p.runs:
        run.text = ''
    if p.runs:
        p.runs[0].text = new_text
    else:
        p.add_run(new_text)
    return True

def delete_para(doc, idx):
    if idx < 0 or idx >= len(doc.paragraphs):
        return
    p = doc.paragraphs[idx]
    p._element.getparent().remove(p._element)

# ============================================================
# FIX 1: 4.3 业务痛点 - 去掉"明确差距、有针对性地提升"
# ============================================================
idx = find_para_by_contains(doc, '自助评估、明确差距、有针对性地提升')
if idx >= 0:
    new_text = '业务痛点：目前机场只有在正式申报"双碳机场"评价时才能得知自己能否达到目标星级，缺乏过程性反馈。机场需要一个工具，在正式申报前自助评估自身得分情况。'
    replace_para_text(doc, idx, new_text)
    changes_log.append(f'[FIX 1] 4.3业务痛点：去掉"明确差距、有针对性地提升"')
else:
    changes_log.append(f'[SKIP 1] 4.3业务痛点：文本未匹配')

# ============================================================
# FIX 2: 4.3 整改优先级排序 - 删除（不给建议）
# ============================================================
idx = find_para_by_contains(doc, '整改优先级排序')
if idx >= 0:
    delete_para(doc, idx)
    changes_log.append(f'[FIX 2] 4.3删除：整改优先级排序')
else:
    changes_log.append(f'[SKIP 2] 整改优先级排序：未找到')

# ============================================================
# FIX 3: 4.3 模块协同 - 添加"仅做打分功能，不给建议内容"
# ============================================================
idx = find_para_by_contains(doc, '本模块为机场账户登录后的完整版')
if idx >= 0:
    text = doc.paragraphs[idx].text
    if '仅做打分功能' not in text:
        new_text = text + ' 本模块仅做打分功能，不给建议内容。'
        replace_para_text(doc, idx, new_text)
        changes_log.append(f'[FIX 3] 4.3模块协同：添加"仅做打分功能，不给建议内容"')
    else:
        changes_log.append(f'[SKIP 3] 4.3模块协同：已包含')
else:
    changes_log.append(f'[SKIP 3] 4.3模块协同：未找到')

# ============================================================
# FIX 4: 4.2 模块协同 - 改为通过算法和AI进行技术标注与匹配
# ============================================================
idx = find_para_by_contains(doc, '复用数据模块的画像引擎做匹配')
if idx >= 0:
    new_text = '模块协同：复用资源库模块的"技术产品库"作为数据底座，通过算法和AI进行技术标注与匹配。'
    replace_para_text(doc, idx, new_text)
    changes_log.append(f'[FIX 4] 4.2模块协同：改为通过算法和AI进行技术标注与匹配')
else:
    changes_log.append(f'[SKIP 4] 4.2模块协同：未找到')

# ============================================================
# FIX 5: 4.9 模块协同 - 去掉4.6和4.7引用
# ============================================================
idx = find_para_by_contains(doc, '4.6 互助平台')
if idx >= 0:
    text = doc.paragraphs[idx].text
    new_text = text.replace('4.6 互助平台（答疑参与）、4.7 培训体系（培训参与）', '新技术推广（技术匹配）')
    replace_para_text(doc, idx, new_text)
    changes_log.append(f'[FIX 5] 4.9模块协同：去掉4.6和4.7引用')
else:
    changes_log.append(f'[SKIP 5] 4.9模块协同：未找到')

# ============================================================
# FIX 6: 1.1 项目背景 - 添加协会数据收集说明
# ============================================================
idx = find_para_by_contains(doc, '二期建设的核心使命，是把协会已经在收集的丰富数据真正')
if idx >= 0:
    text = doc.paragraphs[idx].text
    if '协会对于机场减排数据收集暂没计划' not in text:
        new_text = text + ' 注：协会对于机场减排数据收集暂没计划，二期数据主要基于一期已收集数据及PDF报告抓取。'
        replace_para_text(doc, idx, new_text)
        changes_log.append(f'[FIX 6] 1.1项目背景：添加协会数据收集说明')
    else:
        changes_log.append(f'[SKIP 6] 1.1项目背景：已包含')
else:
    changes_log.append(f'[SKIP 6] 1.1项目背景：未找到')

# ============================================================
# FIX 7: 修改表格 - 去掉"4.7 培训与认证"
# ============================================================
table_fix_count = 0
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if '4.7 培训与认证' in cell.text:
                cell.text = cell.text.replace('4.7 培训与认证', '')
                table_fix_count += 1
            if '4.7培训' in cell.text:
                cell.text = cell.text.replace('4.7培训', '')
                table_fix_count += 1
if table_fix_count > 0:
    changes_log.append(f'[FIX 7] 表格：去掉"4.7 培训与认证"({table_fix_count}处)')
else:
    changes_log.append(f'[SKIP 7] 表格：未找到"4.7 培训与认证"')

# ============================================================
# Save
# ============================================================
doc.save(doc_path)

print("修复完成！")
for log in changes_log:
    print(f"  {log}")
print(f"\n共执行 {len([l for l in changes_log if l.startswith('[FIX')])} 项修复")
