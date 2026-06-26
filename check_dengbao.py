import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

doc_path = r'c:\Users\adone\OneDrive\桌面\AI练习\wendang\fangan\【协会反馈】绿色机场平台二期-建设方案V.020260608_修改版.docx'
doc = Document(doc_path)

print("=" * 60)
print("检查文档中的'等保'相关内容")
print("=" * 60)

# 检查段落
has_dengbao = False
for i, p in enumerate(doc.paragraphs):
    if '等保' in p.text:
        print(f"[段落 {i}] {p.text[:150]}")
        has_dengbao = True

# 检查表格
for table_idx, table in enumerate(doc.tables):
    for row_idx, row in enumerate(table.rows):
        for cell_idx, cell in enumerate(row.cells):
            if '等保' in cell.text:
                print(f"[表格 {table_idx} 行{row_idx} 列{cell_idx}] {cell.text[:150]}")
                has_dengbao = True

if not has_dengbao:
    print("✓ 文档中已无'等保'相关内容")
else:
    print(f"\n⚠ 发现 {sum(1 for p in doc.paragraphs if '等保' in p.text)} 个段落和 {sum(1 for table in doc.tables for row in table.rows for cell in row.cells if '等保' in cell.text)} 个表格单元格包含'等保'")

print("=" * 60)