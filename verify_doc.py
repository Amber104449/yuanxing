import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

doc_path = r'c:\Users\adone\OneDrive\桌面\AI练习\wendang\fangan\【协会反馈】绿色机场平台二期-建设方案V.020260608_修改版.docx'
doc = Document(doc_path)

with open(r'c:\Users\adone\OneDrive\桌面\AI练习\wendang\fangan\modified_content.txt', 'w', encoding='utf-8') as f:
    f.write("=== PARAGRAPHS ===\n")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            f.write(f'[{i}] {p.style.name}: {p.text}\n')
    
    f.write('\n=== TABLES ===\n')
    for ti, table in enumerate(doc.tables):
        f.write(f'\nTable {ti}:\n')
        for ri, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells]
            f.write(f'  Row {ri}: {cells}\n')

print("Done")
