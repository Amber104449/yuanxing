import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

doc_path = r'c:\Users\adone\OneDrive\桌面\AI练习\wendang\fangan\【协会反馈】绿色机场平台二期-建设方案V.020260608_修改版.docx'
doc = Document(doc_path)

changes_log = []

# Helper functions
def find_para_idx(doc, text_fragment):
    for i, p in enumerate(doc.paragraphs):
        if text_fragment in p.text:
            return i
    return -1

def replace_para_text(doc, idx, new_text):
    if idx < 0 or idx >= len(doc.paragraphs):
        return False
    p = doc.paragraphs[idx]
    for run in p.runs:
        run.text = ''
    if p.runs:
        p.runs[0].text = new_text
    else:
        p.add_run(new_text)
    return True

def delete_para(doc, idx):
    if idx < 0 or idx >= len(doc.paragraphs):
        return
    p = doc.paragraphs[idx]
    p._element.getparent().remove(p._element)

# ============================================================
# FIX 1: 4.3 业务痛点 - 去掉"明确差距、有针对性地提升"
# ============================================================
idx = find_para_idx(doc, '业务痛点：目前机场只有在正式申报"双碳机场"评价时才能得知自己能否达到目标星级')
if idx >= 0:
    text = doc.paragraphs[idx].text
    if '明确差距' in text or '有针对性地提升' in text:
        new_text = '业务痛点：目前机场只有在正式申报"双碳机场"评价时才能得知自己能否达到目标星级，缺乏过程性反馈。机场需要一个工具，在正式申报前自助评估自身得分情况。'
        replace_para_text(doc, idx, new_text)
        changes_log.append(f'[FIX 4.3] 业务痛点：去掉"明确差距、有针对性地提升"')

# ============================================================
# FIX 2: 4.3 整改优先级排序 - 删除（不给建议）
# ============================================================
idx = find_para_idx(doc, '整改优先级排序：结合"提升性价比"')
if idx >= 0:
    delete_para(doc, idx)
    changes_log.append(f'[FIX 4.3] 删除：整改优先级排序（不给建议）')

# ============================================================
# FIX 3: 4.2 模块协同 - 改为通过算法和AI进行技术标注与匹配
# ============================================================
idx = find_para_idx(doc, '模块协同：复用资源库模块的"技术产品库"作为数据底座，复用数据模块的画像引擎做匹配。')
if idx >= 0:
    new_text = '模块协同：复用资源库模块的"技术产品库"作为数据底座，通过算法和AI进行技术标注与匹配。'
    replace_para_text(doc, idx, new_text)
    changes_log.append(f'[FIX 4.2] 模块协同：改为通过算法和AI进行技术标注与匹配')

# ============================================================
# FIX 4: 4.11 协会业务扩展模块总览 - 更新模块数量
# ============================================================
idx = find_para_idx(doc, '9 项扩展模块按服务对象与服务目的可归为四组')
if idx >= 0:
    new_text = '多项扩展模块按服务对象与服务目的可归为四组：'
    replace_para_text(doc, idx, new_text)
    changes_log.append(f'[FIX 4.11] 修改：9项→多项')

idx = find_para_idx(doc, '核心理念：二期平台不再是"协会用来开展评价工作的内部系统"')
if idx >= 0:
    text = doc.paragraphs[idx].text
    if '9 项扩展模块' in text:
        new_text = text.replace('9 项扩展模块', '多项扩展模块')
        replace_para_text(doc, idx, new_text)
        changes_log.append(f'[FIX 4.11] 核心理念：9项→多项')

# ============================================================
# FIX 5: 4.3 模块协同 - 添加"仅做打分功能，不给建议内容"
# ============================================================
idx = find_para_idx(doc, '模块协同：复用第三章的"评价指标失分热力图"分析引擎；与插件工具模块的"机场自评分速算工具"差异')
if idx >= 0:
    text = doc.paragraphs[idx].text
    if '仅做打分功能' not in text:
        new_text = text + '本模块仅做打分功能，不给建议内容。'
        replace_para_text(doc, idx, new_text)
        changes_log.append(f'[FIX 4.3] 模块协同：添加"仅做打分功能，不给建议内容"')

# ============================================================
# FIX 6: 修改表格6 - 去掉"4.7 培训与认证"
# ============================================================
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if '4.7 培训与认证' in cell.text:
                cell.text = cell.text.replace('4.7 培训与认证', '')
                changes_log.append(f'[FIX Table6] 去掉"4.7 培训与认证"')
            if '4.6 互助平台' in cell.text:
                cell.text = cell.text.replace('4.6 互助平台', '互助协作')
                changes_log.append(f'[FIX Table] 4.6互助平台→互助协作')

# ============================================================
# FIX 7: 修改1.1 项目背景 - 添加说明
# ============================================================
idx = find_para_idx(doc, '二期建设的核心使命，是把协会已经在收集的丰富数据真正"激活"')
if idx >= 0:
    text = doc.paragraphs[idx].text
    if '协会对于机场减排数据收集暂没计划' not in text:
        new_text = text + '注：协会对于机场减排数据收集暂没计划，二期数据主要基于一期已收集数据及PDF报告抓取。'
        replace_para_text(doc, idx, new_text)
        changes_log.append(f'[FIX 1.1] 添加：协会对于机场减排数据收集暂没计划的说明')

# ============================================================
# Save
# ============================================================
doc.save(doc_path)

print("修复完成！")
for log in changes_log:
    print(f"  {log}")
print(f"\n共执行 {len(changes_log)} 项修复")
