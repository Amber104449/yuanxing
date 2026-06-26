import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

doc_path = r'c:\Users\adone\OneDrive\桌面\AI练习\wendang\fangan\【协会反馈】绿色机场平台二期-建设方案V.020260608_最终修改版.docx'
doc = Document(doc_path)

def find_para_by_contains(doc, text_fragment):
    for i, p in enumerate(doc.paragraphs):
        if text_fragment in p.text:
            return i
    return -1

# 检查段落[202]
idx = 202
if idx < len(doc.paragraphs):
    text = doc.paragraphs[idx].text
    print(f"段落[202]当前内容: {text}")

# 修复"9 项扩展模块" -> "7项扩展模块"
idx = find_para_by_contains(doc, '9 项扩展模块按服务对象与服务目的可归为四组')
if idx >= 0:
    text = doc.paragraphs[idx].text
    new_text = text.replace('9 项扩展模块', '7项扩展模块')
    # 替换
    p = doc.paragraphs[idx]
    for run in p.runs:
        run.text = ''
    p.runs[0].text = new_text
    print(f"[修复] 段落{idx}已更新")
else:
    print("[修复] 未找到需要修复的段落")

doc.save(doc_path)
print("\n文档已保存")

# 最终验证
doc2 = Document(doc_path)
print("\n=== 最终验证 ===")
for i, p in enumerate(doc2.paragraphs):
    if '9' in p.text and '项扩展模块' in p.text:
        print(f"⚠ 段落[{i}]仍包含9项: {p.text}")
    if '7项' in p.text and '扩展模块' in p.text:
        print(f"✅ 段落[{i}]已更新为7项: {p.text}")