import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

doc_path = r'c:\Users\adone\OneDrive\桌面\AI练习\wendang\fangan\【协会反馈】绿色机场平台二期-建设方案V.020260608_最终修改版.docx'
doc = Document(doc_path)

print("=" * 70)
print("最终修改版文档验证")
print("=" * 70)

# 检查1: 等保三级是否已删除
dengbao_count = 0
for p in doc.paragraphs:
    if '等保' in p.text:
        dengbao_count += 1
        print(f"[等保] 段落: {p.text[:80]}")
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if '等保' in cell.text:
                dengbao_count += 1
                print(f"[等保] 表格: {cell.text[:80]}")
print(f"\n等保相关内容数量: {dengbao_count} (应为0)")

# 检查2: 7项增值服务
print("\n" + "-" * 70)
print("增值服务相关段落:")
for i, p in enumerate(doc.paragraphs):
    if '7项' in p.text or '9类' in p.text or '9 项' in p.text:
        print(f"[{i}] {p.text}")

# 检查3: 图片转文字
print("\n" + "-" * 70)
print("PDF抓取方式相关段落:")
for i, p in enumerate(doc.paragraphs):
    if '图片转文字' in p.text or '避免AI' in p.text:
        print(f"[{i}] {p.text}")

# 检查4: 验证关键位置的具体文本
print("\n" + "-" * 70)
print("关键位置验证:")

# 1.2目标二
for i, p in enumerate(doc.paragraphs):
    if '目标二' in p.text and '增值服务' in p.text:
        print(f"[1.2目标二] {p.text}")

# 4.1
for i, p in enumerate(doc.paragraphs):
    if '4.1' in p.text or '设计动机' in p.text:
        if '7项' in p.text or '多项' in p.text:
            print(f"[4.1] {p.text[:100]}")

# 4.11
for i, p in enumerate(doc.paragraphs):
    if '4.11' in p.text or '扩展模块' in p.text:
        if '7项' in p.text or '9项' in p.text:
            print(f"[4.11] {p.text}")

print("\n" + "=" * 70)
if dengbao_count == 0:
    print("✅ 等保内容已删除")
else:
    print("❌ 等保内容仍存在")

if any('7项' in p.text for p in doc.paragraphs if '目标二' in p.text):
    print("✅ 7项增值服务已更新")
else:
    print("❌ 7项增值服务未正确更新")

if any('图片转文字' in p.text for p in doc.paragraphs):
    print("✅ 图片转文字算法已添加")
else:
    print("❌ 图片转文字算法未添加")

print("=" * 70)