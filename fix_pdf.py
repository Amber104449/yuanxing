import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

doc_path = r'c:\Users\adone\OneDrive\桌面\AI练习\wendang\fangan\【协会反馈】绿色机场平台二期-建设方案V.020260608_最终修改版.docx'
doc = Document(doc_path)

# 找到段落
idx = -1
for i, p in enumerate(doc.paragraphs):
    if '四星评价强制' in p.text and '碳达峰路径规划研究' in p.text:
        idx = i
        break

if idx >= 0:
    text = doc.paragraphs[idx].text
    # 添加PDF抓取方式的修改内容
    if '图片转文字' not in text:
        new_text = text + ' 二期新建立数据库，通过图片转文字算法逻辑做抓取，避免AI的直接参与导致数据信息泄露。'
        # 替换段落内容
        p = doc.paragraphs[idx]
        for run in p.runs:
            run.text = ''
        p.runs[0].text = new_text
        print(f"[修改3] 段落{idx}已更新：添加图片转文字算法逻辑")
    else:
        print(f"[修改3] 段落{idx}已包含图片转文字内容")
else:
    print("[修改3] 未找到目标段落")

# 保存文档
doc.save(doc_path)
print("文档已保存")